"""
tests/test_presencia_presentacion.py — Presentación de la presencia parcial
(§1.8 de REGLAS.md) fuera del motor: dashboard, estadísticas, equidad y
exportación a Excel (paso 5).

El contrato de presencia parcial produce días en estado "Fuera" que "no
computan en nada" (ni libres, ni equidad, ni cobertura, ni festivos). Este
módulo verifica las cuatro capas de presentación sobre ese principio:

  1. Dashboard: un día "Fuera" se trata como cualquier día sin trabajar
     (mismo camino que Libre/Vacaciones): sin horario, sin compañeros y con
     "vuelves el…" calculado contra el próximo turno real.
  2. Estadísticas: "Fuera" no es un turno contable — ni en los totales
     históricos ni en el resumen por periodo.
  3. Equidad: un rol con días de presencia queda fuera de la media del
     equipo de forma ESTRUCTURAL, aunque su excluir_de_equidad llegue sin
     marcar (roles.json editado a mano o anterior a esta regla).
  4. Excel: la celda "Fuera" se exporta con su gris propio (SHIFT_COLORS) y
     no tiene fila en COBERTURA (no es un turno).
"""

from datetime import date, timedelta

import pytest

SEMANA_INICIO = date(2099, 3, 2)  # Lunes en el futuro lejano


# ─── Fixtures compartidas ────────────────────────────────────────────────────

@pytest.fixture
def data_dir(tmp_path, monkeypatch):
    backup = tmp_path / "backups"
    backup.mkdir()

    import herramientas
    import cupos
    import app as app_module

    monkeypatch.setattr(herramientas, "DATA_DIR", str(tmp_path))
    monkeypatch.setattr(herramientas, "BACKUP_DIR", str(backup))
    monkeypatch.setattr(herramientas, "INTERVENCIONES_FILE", str(tmp_path / "intervenciones.json"))
    monkeypatch.setattr(cupos, "CUPOS_EXTRA_FILE", str(tmp_path / "cupos_extra.json"))
    monkeypatch.setattr(app_module, "DATA_DIR", str(tmp_path))
    return tmp_path


@pytest.fixture
def client(data_dir):
    import app as app_module
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client()


def _sesion(client, *, uid, username, rol, nombre, vinculado):
    with client.session_transaction() as sess:
        sess["uid"] = uid
        sess["username"] = username
        sess["rol"] = rol
        sess["nombre"] = nombre
        sess["vinculado"] = vinculado


def _lunes_actual() -> date:
    hoy = date.today()
    return hoy - timedelta(days=hoy.weekday())


def _seed_plantilla_roles(roles: dict, dept_id="recepcion"):
    """roles: {nombre: role}."""
    from herramientas import guardar_plantilla
    guardar_plantilla(
        [{"name": n, "role": r, "preferences": []} for n, r in roles.items()],
        dept_id,
    )


def _seed_semana(lunes, schedule, roles=None, dept_id="recepcion"):
    """Crea/actualiza una semana de historial: schedule {nombre: [7 turnos]};
    la instantánea de staff guarda el rol indicado (por defecto "normal")."""
    from herramientas import cargar_historial, guardar_historial, _semana_str
    roles = roles or {}
    historial = cargar_historial(dept_id)
    historial[_semana_str(lunes)] = {
        "schedule": schedule,
        "staff": [{"name": n, "role": roles.get(n, "normal"), "preferences": []}
                  for n in schedule],
    }
    guardar_historial(historial, dept_id)


_ROL_BASE = dict(
    label="Test", icon="", turno_fijo=None, dias_libres_fijos=None,
    dias_libres_sustituto=None, dias_presencia=[], sin_noches=False,
    unico=False, protegido=False, borrable=True, cubre_ausencia_de=None,
    cobertura_semana_completa=False, cubrir_dias_libres_fijos=False,
    excluir_de_equidad=False,
)


# ═══════════════════════════════════════════════════════════════════════════
# 1. Dashboard: "Fuera" es un día sin trabajar
# ═══════════════════════════════════════════════════════════════════════════

def test_dashboard_fuera_dia_sin_trabajar_con_regreso_y_sin_companeros(client, data_dir):
    """Si "hoy" cae en un día "Fuera", el dashboard lo presenta como
    cualquier día sin trabajar: sin horario ni compañeros de turno, y
    "vuelves el…" apunta al primer turno real saltando los días fuera."""
    lunes = _lunes_actual()
    _seed_plantilla_roles({"Ref": "refuerzo", "Ana": "normal"})
    # Toda la semana actual en "Fuera" para ambos; la siguiente arranca con
    # Mañana. Así el primer turno real tras hoy es siempre el Lunes
    # siguiente, sea cual sea el día en que corra la suite.
    _seed_semana(lunes, {"Ref": ["Fuera"] * 7, "Ana": ["Fuera"] * 7})
    _seed_semana(lunes + timedelta(weeks=1),
                 {"Ref": ["Mañana"] * 7, "Ana": ["Mañana"] * 7})

    _sesion(client, uid="1", username="ref", rol="empleado",
            nombre="Ref", vinculado="Ref")
    r = client.get("/api/dashboard")
    assert r.status_code == 200
    data = r.get_json()["data"]

    ahora = data["ahora_mismo"]
    assert ahora["turno_hoy"] == "Fuera"
    assert ahora["horario_hoy"] is None, "un día fuera no tiene tramos horarios"
    assert ahora["companeros_hoy"] == [], (
        "Ana también está 'Fuera' hoy: no es una compañera de turno"
    )
    assert ahora["regreso"] == {
        "fecha": (lunes + timedelta(days=7)).isoformat(),
        "dia": "Lunes", "turno": "Mañana",
    }, "el regreso salta toda la semana fuera y apunta al primer turno real"

    # La tira semanal muestra los siete días "Fuera" tal cual; el regreso
    # cae fuera de la semana actual, así que ningún día va marcado (↩).
    assert [d["turno"] for d in data["semana"]["dias"]] == ["Fuera"] * 7
    assert not any(d["es_regreso"] for d in data["semana"]["dias"])


def test_dashboard_marca_regreso_en_la_tira_si_el_turno_real_cae_esta_semana(client, data_dir):
    """Si tras un día "Fuera" hay un turno real más adelante en la misma
    semana, ese día queda marcado como regreso (↩) en la tira semanal."""
    hoy_idx = date.today().weekday()
    lunes = _lunes_actual()

    turnos_ref = ["Fuera"] * 7
    if hoy_idx < 6:
        # Deja un turno real el día siguiente a hoy, dentro de la semana.
        turnos_ref[hoy_idx + 1] = "Tarde"
        fecha_regreso = lunes + timedelta(days=hoy_idx + 1)
        turno_regreso = "Tarde"
    else:
        # Domingo: el primer turno real ya cae en la semana siguiente.
        fecha_regreso = lunes + timedelta(days=7)
        turno_regreso = "Mañana"

    _seed_plantilla_roles({"Ref": "refuerzo"})
    _seed_semana(lunes, {"Ref": turnos_ref})
    _seed_semana(lunes + timedelta(weeks=1), {"Ref": ["Mañana"] * 7})

    _sesion(client, uid="1", username="ref", rol="empleado",
            nombre="Ref", vinculado="Ref")
    data = client.get("/api/dashboard").get_json()["data"]

    assert data["ahora_mismo"]["turno_hoy"] == "Fuera"
    assert data["ahora_mismo"]["regreso"]["fecha"] == fecha_regreso.isoformat()
    assert data["ahora_mismo"]["regreso"]["turno"] == turno_regreso
    for d in data["semana"]["dias"]:
        assert d["es_regreso"] == (d["fecha"] == fecha_regreso.isoformat()), (
            f"{d['fecha']}: es_regreso debe marcar exactamente el día del regreso"
        )


# ═══════════════════════════════════════════════════════════════════════════
# 2. Estadísticas: "Fuera" no computa
# ═══════════════════════════════════════════════════════════════════════════

def test_estadisticas_fuera_no_cuenta_como_turno(client, data_dir):
    """'Fuera' no es una clave de turnos: no aparece en los contadores
    históricos ni en el resumen por periodo, y no altera los reales."""
    from estadisticas import estadisticas_de, resumen_periodo

    lunes = _lunes_actual()
    _seed_plantilla_roles({"Ref": "refuerzo"})
    for i in range(2):
        _seed_semana(lunes + timedelta(weeks=i),
                     {"Ref": ["Mañana"] * 5 + ["Fuera", "Fuera"]})

    e = estadisticas_de("Ref")
    assert e["turnos"] == {"Mañana": 10, "Tarde": 0, "Noche": 0,
                           "Partido": 0, "Libre": 0}, (
        "solo los turnos reales cuentan; el finde 'Fuera' no suma a Libre"
    )

    r = resumen_periodo("Ref", periodo="3m")
    assert r["turnos"]["Mañana"] == 10
    assert "Fuera" not in r["turnos"]


# ═══════════════════════════════════════════════════════════════════════════
# 3. Equidad: exclusión estructural por presencia
# ═══════════════════════════════════════════════════════════════════════════

def test_equidad_excluye_rol_con_presencia_aunque_el_flag_llegue_sin_marcar(client, data_dir):
    """Un rol con dias_presencia queda fuera de la media del equipo por ser
    presencia en sí (estadisticas._calcular_equidad), no solo por el flag
    excluir_de_equidad: un roles.json editado a mano que no lo marque no
    debe colar un refuerzo de 2 días en el balanceo de findes."""
    from herramientas import guardar_roles

    guardar_roles({
        "normal":   {**_ROL_BASE, "label": "Empleado", "protegido": True,
                     "borrable": False},
        "refuerzo": {**_ROL_BASE, "label": "Refuerzo finde",
                     "dias_presencia": [5, 6],
                     "excluir_de_equidad": False},  # a propósito sin marcar
    }, "recepcion")

    roles = {"Ana": "normal", "Luis": "normal", "Ref": "refuerzo"}
    _seed_plantilla_roles(roles)

    # 20 semanas: Ana y Luis libran el finde 4 veces cada una (ratio 0.2);
    # Ref lo libra las 20 (ratio 1.0). Si colara en la media, la desplazaría
    # de 0.2 a ~0.467.
    personas_libre = {"Ana": {0, 5, 10, 15}, "Luis": {2, 7, 12, 17},
                      "Ref": set(range(20))}
    for i in range(20):
        lunes = SEMANA_INICIO + timedelta(weeks=i)
        libran = {n for n, semanas in personas_libre.items() if i in semanas}
        _seed_semana(lunes, {
            n: (["Mañana"] * 5 + ["Libre", "Libre"]) if n in libran else ["Mañana"] * 7
            for n in roles
        }, roles=roles)

    _sesion(client, uid="1", username="ana", rol="empleado",
            nombre="Ana", vinculado="Ana")
    data = client.get("/api/dashboard").get_json()["data"]
    assert data["equidad"]["findes"]["media_equipo"] == pytest.approx(0.2, abs=1e-6), (
        "el refuerzo con presencia no debe entrar en la media del equipo"
    )

    from estadisticas import equidad_de
    eq_ref = equidad_de("Ref")
    assert eq_ref["excluido_de_balanceo"] is True, (
        "su bloque de equidad propio se marca como excluido de balanceo"
    )


# ═══════════════════════════════════════════════════════════════════════════
# 4. Excel: celda "Fuera" con color propio y sin fila de cobertura
# ═══════════════════════════════════════════════════════════════════════════

def test_excel_celda_fuera_con_color_propio_y_sin_fila_de_cobertura(tmp_path):
    """La celda 'Fuera' se exporta con su gris (SHIFT_COLORS['Fuera']), los
    turnos reales conservan su color, y las filas de COBERTURA siguen siendo
    las cinco de siempre — el COUNTIF por texto exacto nunca cuenta 'Fuera'."""
    import copy
    from modelo_base import export_to_excel, SHIFT_COLORS
    from modelo_recepcion import RECEPCION_CONFIG
    from openpyxl import load_workbook

    cfg = copy.deepcopy(RECEPCION_CONFIG)
    staff = [
        {"name": "Ref", "role": "refuerzo"},
        {"name": "Ana", "role": "normal"},
    ]
    schedule = {
        "Ref": ["Fuera"] * 5 + ["Mañana", "Mañana"],
        "Ana": ["Mañana"] * 7,
    }
    ruta = tmp_path / "export_fuera.xlsx"
    export_to_excel(schedule, staff, SEMANA_INICIO, str(ruta), cfg)

    wb = load_workbook(str(ruta))
    ws = wb.active

    # Ref es la primera persona (fila 3; 1=título, 2=cabecera de días).
    fila_ref = 3
    assert ws.cell(row=fila_ref, column=2).value == "Fuera"
    assert ws.cell(row=fila_ref, column=2).fill.start_color.rgb[-6:] == SHIFT_COLORS["Fuera"]
    assert ws.cell(row=fila_ref, column=6).fill.start_color.rgb[-6:] == SHIFT_COLORS["Fuera"]
    assert ws.cell(row=fila_ref, column=7).value == "Mañana"
    assert ws.cell(row=fila_ref, column=7).fill.start_color.rgb[-6:] == SHIFT_COLORS["Mañana"]

    # Filas de COBERTURA: las cinco de siempre; ninguna para "Fuera".
    etiquetas = [ws.cell(row=r, column=1).value for r in range(1, 20)]
    assert "COBERTURA" in etiquetas
    assert "Fuera" not in etiquetas
